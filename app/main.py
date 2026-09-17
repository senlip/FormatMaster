"""FormatMaster 程序入口。"""
from __future__ import annotations

import os
import sys
from pathlib import Path

# 允许以源码方式直接运行（python app/main.py）
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from PySide6.QtCore import Qt  # noqa: E402
from PySide6.QtGui import QIcon  # noqa: E402
from PySide6.QtWidgets import QApplication, QMessageBox  # noqa: E402


def _kugou_fingerprint() -> str:
    """酷狗密钥流指纹 —— 用来确认打包没漏掉 ``MASK_V2_PRE_DEF`` 那一层。

    密钥与密文都置零时，酷狗的解密公式退化成"输出 = 本次生效的密钥流"，
    所以这一个值就能和 ``tests/fixtures`` 里的官方金标直接对照。
    （酷狗曾经因为漏掉一层异或，99% 的字节都解错，而自测还全绿。）
    """
    try:
        from app.core.decryptors.kugou import decrypt_bytes

        return decrypt_bytes(bytes(4096), bytes(17), False)[:16].hex()
    except Exception as exc:                       # 缺资源等，报告里如实写出来
        return f"error: {exc}"


def _document_smoke_test() -> dict:
    """文档转换冒烟测试：真的转一份带中文的 docx，并检查产物内容。

    为什么要放在打包自检里：DOCX → PDF 曾经在冻结环境下把 DOCX 的 ZIP 流
    当纯文本读，排出几十页乱码还报"成功"。只探测依赖是否 import 得到
    （见 engine_health）根本发现不了这类问题 —— **必须看产物内容**。

    检查项与 `tests/selftest_docs.py` 一致：文本层含已知中文串、中文字体
    真的嵌入（``ext != 'n/a'``）、页数合理。
    """
    import tempfile

    from app.core.engines.base import ConvertOptions
    from app.core.engines.doc_engine import DocumentEngine

    mark = "氧化铝沉降槽底流密度在线检测"
    work = Path(tempfile.gettempdir()) / "formatmaster_doc_smoke"
    work.mkdir(parents=True, exist_ok=True)
    src = work / "smoke.docx"
    report: dict = {}

    try:
        from docx import Document as DocxDocument

        document = DocxDocument()
        document.add_heading(mark, level=1)
        document.add_paragraph("拜耳法工艺参数：ρ = 1450 kg/m³。")
        document.save(src)
    except Exception as exc:
        return {"ok": False, "stage": "make_sample", "error": str(exc)}

    engine = DocumentEngine()
    for mode in ("builtin", "auto"):
        entry: dict = {}
        dst = work / f"smoke_{mode}.pdf"
        try:
            result = engine.convert(
                src, dst,
                ConvertOptions(target_ext="pdf", overwrite=True, doc_engine=mode),
            )
            entry["ok"] = bool(result.ok)
            entry["message"] = result.message[:200]
            if result.ok and dst.exists():
                import pymupdf as fitz

                with fitz.open(dst) as doc:
                    text = "\n".join(p.get_text() for p in doc)
                    fonts = {f[3]: f[1] for p in doc for f in p.get_fonts(full=True)}
                    entry.update({
                        "bytes": dst.stat().st_size,
                        "pages": doc.page_count,
                        "chars": len(text),
                        "mark_found": mark in text,
                        "fonts": sorted(fonts),
                        "unembedded": [n for n, ext in fonts.items() if ext == "n/a"],
                    })
                    entry["ok"] = bool(
                        result.ok and entry["mark_found"]
                        and not entry["unembedded"] and doc.page_count <= 40
                    )
        except Exception as exc:
            entry = {"ok": False, "message": f"{type(exc).__name__}: {exc}"[:200]}
        report[mode] = entry

    report["ok"] = all(item.get("ok") for item in report.values()
                       if isinstance(item, dict))
    return report


def _selftest_report() -> int:
    """打包后自检：确认引擎二进制与资源路径在冻结环境下依然可用。

    用环境变量 FORMATMASTER_SELFTEST=1 触发，结果写到系统临时目录的
    selftest.json，便于排查"源码能跑、打包后找不到 ffmpeg / 转出乱码"这类问题。
    """
    import json
    import tempfile
    import time

    from app.core.decryptors import sanity_check
    from app.core.registry import get_registry
    from app.core.runtime import PROJECT_ROOT as ROOT
    from app.core.runtime import environment_report

    registry = get_registry()
    health = registry.health()             # 触发各引擎的可用性探测
    report = {
        "generated_at": time.time(),
        "frozen": bool(getattr(sys, "frozen", False)),
        "meipass": getattr(sys, "_MEIPASS", None),
        "executable": sys.executable,
        "project_root": str(ROOT),
        "engines": environment_report(),
        "engine_health": health,
        "decryptors": sanity_check(),
        "kugou_keystream_fingerprint": _kugou_fingerprint(),
        "document_smoke": _document_smoke_test(),
        "routes": {
            ext: [spec.ext for spec in registry.targets_for(ext)]
            for ext in ("mp4", "mp3", "png", "pdf", "docx", "zip", "txt")
        },
        "encrypted_routes": {
            ext: [spec.ext for spec in registry.targets_for(ext)]
            for ext in ("ncm", "qmcflac", "kgm", "kgma", "vpr", "kwm")
        },
    }
    # 写到系统临时目录：不污染交付目录，也避免打包流程里出现删除动作
    target = Path(tempfile.gettempdir()) / "formatmaster_selftest.json"
    target.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"自检报告已写入 {target}")
    smoke = report["document_smoke"]
    print(f"文档冒烟测试：{'通过' if smoke.get('ok') else '未通过'}")
    return 0 if smoke.get("ok") else 1


def _icon_path() -> Path | None:
    for name in ("app.ico", "app.png"):
        candidate = Path(__file__).resolve().parent / "assets" / name
        if candidate.is_file():
            return candidate
    return None


def main() -> int:
    if os.environ.get("FORMATMASTER_SELFTEST"):
        return _selftest_report()

    QApplication.setHighDpiScaleFactorRoundingPolicy(
        Qt.HighDpiScaleFactorRoundingPolicy.PassThrough
    )
    app = QApplication(sys.argv)
    app.setApplicationName("FormatMaster")
    app.setApplicationDisplayName("FormatMaster 万能格式转换器")
    app.setOrganizationName("ErBai")

    from app.ui import theme

    app.setStyleSheet(theme.stylesheet())        # 必须在 QApplication 建好之后
    font = app.font()
    font.setPointSize(9)
    app.setFont(font)

    icon = _icon_path()
    if icon:
        app.setWindowIcon(QIcon(str(icon)))

    try:
        from app.ui.main_window import MainWindow
    except Exception as exc:                       # 依赖缺失时给出人话提示
        QMessageBox.critical(
            None, "启动失败",
            f"程序初始化失败：\n{exc}\n\n请确认已安装 requirements.txt 中的依赖。",
        )
        return 1

    window = MainWindow()
    if icon:
        window.setWindowIcon(QIcon(str(icon)))
    window.show()
    return app.exec()


if __name__ == "__main__":
    sys.exit(main())
